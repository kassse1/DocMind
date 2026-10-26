"""Build dogovor_arendy_sklada.docx (the file uploaded in the scenario run) from its text.

    python docs/evaluation/scenarios/make_docx.py

Paragraphs are separated by blank lines; the first becomes a level-1 heading.
"""

from pathlib import Path

from docx import Document

here = Path(__file__).parent
text = (here / "dogovor_arendy_sklada.txt").read_text(encoding="utf-8")
paragraphs = [p.strip() for p in text.split("\n\n") if p.strip()]
document = Document()
document.add_heading(paragraphs[0], 1)
for paragraph in paragraphs[1:]:
    document.add_paragraph(paragraph)
document.save(here / "dogovor_arendy_sklada.docx")

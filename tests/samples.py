"""Build small PDF/DOCX/ZIP files in memory, so no binary fixtures live in the repo."""

import io
import zipfile

import docx
import pymupdf


def make_pdf(pages: list[str], *, user_password: str | None = None) -> bytes:
    """Base-14 fonts only cover Latin-1, so PDF samples use ASCII text."""
    document = pymupdf.open()
    for text in pages:
        page = document.new_page()
        if text:
            page.insert_text((72, 72), text)
    if user_password is None:
        return document.tobytes()
    return document.tobytes(
        encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw=user_password, owner_pw="owner"
    )


def make_docx(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    document = docx.Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                grid.cell(r, c).text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_zip(members: dict[str, bytes]) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in members.items():
            archive.writestr(name, data)
    return buffer.getvalue()
